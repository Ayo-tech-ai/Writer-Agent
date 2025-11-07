# =====================================================================
# 🚀 ENHANCED STREAMLIT AGENTIC RESEARCH APP (GROQ + SERPER API)
# =====================================================================

import os
import streamlit as st
import requests
import json
from datetime import datetime
from io import BytesIO
from docx import Document
from functools import lru_cache
import time

# =====================================================================
# ⚙️ APP CONFIGURATION
# =====================================================================

st.set_page_config(
    page_title="LinkedIn Article Generator",
    page_icon="🌍",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- Responsive UI Styling ---
st.markdown("""
    <style>
        .main {
            padding: 1.2rem;
        }
        @media (min-width: 1024px) {
            .main > div {
                max-width: 1000px;
                margin: auto;
            }
        }
        @media (max-width: 768px) {
            .main {
                padding: 0.8rem;
            }
        }
        h1, h2, h3 {
            color: #222;
        }
        .stProgress > div > div > div > div {
            background-color: #4CAF50;
        }
    </style>
""", unsafe_allow_html=True)

st.title("🌍 LinkedIn Article Generator")
st.markdown(
    """
    🤖 Powered by Groq + Serper API  
    Perform intelligent web research with reliable search results and Generate a LinkedIn Post.
    """
)

# =====================================================================
# 🎯 CONSTANTS & PROMPT TEMPLATES
# =====================================================================

LINKEDIN_POST_PROMPT = """
Create a professional LinkedIn post based on this research. Follow these guidelines:

**STRUCTURE:**
🎯 Catchy Headline with Emoji
💡 Strong Opening Hook
📊 Key Insights (3-5 bullet points)
🚀 Practical Implications
🤔 Thought-provoking Question
🏷️ Relevant Hashtags

**TONE:**
- Professional yet conversational
- Data-driven but accessible
- Engaging and shareable
- Authentic voice
- Humanize

**FORMATTING:**
- Use emojis sparingly for visual breaks
- Short paragraphs (2-3 lines max)
- Clear section separation
- Mobile-friendly formatting

Research Content: {research_content}

Generate the LinkedIn post accordingly:
"""

# =====================================================================
# 🔑 API KEY INPUT & VALIDATION
# =====================================================================

def validate_api_keys(groq_key, serper_key):
    """Validate API keys before proceeding"""
    errors = []
    if not groq_key:
        errors.append("❌ Groq API key is required")
    elif not groq_key.startswith(('gsk_', 'gpk_')):
        errors.append("❌ Invalid Groq API key format (should start with gsk_ or gpk_)")
    
    if not serper_key:
        errors.append("❌ Serper API key is required")
    elif len(serper_key) < 10:
        errors.append("❌ Invalid Serper API key")
    
    return errors

with st.sidebar:
    st.header("🔐 Configuration")

    # API Key inputs
    groq_api_key = st.text_input(
        "Enter your Groq API Key",
        type="password",
        placeholder="gsk_... or leave blank to use GROQ_API_KEY env var"
    )

    serper_api_key = st.text_input(
        "Serper API Key (Required)",
        type="password",
        placeholder="Enter your Serper API key",
        help="Get free key from https://serper.dev"
    )

    # Model selection
    model_options = {
        "Llama 3.3 70B Versatile": "llama-3.3-70b-versatile",
        "Llama 3.1 8B Instant": "llama-3.1-8b-instant",
    }
    selected_model = st.selectbox(
        "Select Model",
        options=list(model_options.keys()),
        index=0
    )

    # Configuration options
    col1, col2 = st.columns(2)
    with col1:
        temperature = st.slider("Creativity", 0.0, 1.0, 0.7, 0.1)
    with col2:
        max_results = st.slider("Max Results", 1, 10, 5)

    st.markdown("---")
    st.header("⚙️ Advanced Settings")
    
    # Post customization
    post_tone = st.selectbox(
        "Post Tone",
        ["Professional", "Conversational", "Inspirational", "Technical", "Casual"]
    )
    
    include_statistics = st.checkbox("Include Statistics", value=True)
    add_call_to_action = st.checkbox("Add Call-to-Action", value=True)
    
    # Target audience
    audience_options = ["Executives", "Managers", "Technical", "General", "Students", "Entrepreneurs"]
    audience = st.multiselect(
        "Target Audience",
        audience_options,
        default=["General"]
    )

    st.markdown("---")
    st.info("💡 **Using Serper API for reliable web search**")

# =====================================================================
# 🔍 ENHANCED SERPER API SEARCH IMPLEMENTATION
# =====================================================================

def safe_serper_search(query: str, max_results: int = 5, api_key: str = None) -> str:
    """Enhanced search with better error handling and validation"""
    if not api_key:
        return "❌ Serper API key not provided"
    
    # Query validation
    if not query or len(query.strip()) < 2:
        return "❌ Search query too short"
    
    if len(query) > 300:
        return "❌ Search query too long (max 300 characters)"
    
    try:
        st.write(f"🔍 Searching Serper API for: '{query}'")

        url = "https://google.serper.dev/search"
        payload = json.dumps({
            "q": query,
            "num": max_results,
            "gl": "us",
            "hl": "en"
        })
        headers = {
            'X-API-KEY': api_key,
            'Content-Type': 'application/json'
        }

        with st.spinner("Searching the web..."):
            response = requests.post(url, headers=headers, data=payload, timeout=30)
            response.raise_for_status()

        data = response.json()
        results = []

        # Process organic results
        if 'organic' in data and data['organic']:
            for i, result in enumerate(data['organic'][:max_results], 1):
                title = result.get('title', 'No title')
                link = result.get('link', 'No URL')
                snippet = result.get('snippet', 'No description')
                results.append(f"### 📄 Result {i}: {title}\n**URL:** {link}\n**Summary:** {snippet}\n")
                st.write(f"✅ Found: {title[:80]}...")

        # Also check for news results if organic results are limited
        if len(results) < max_results and 'news' in data and data['news']:
            news_to_add = max_results - len(results)
            for i, result in enumerate(data['news'][:news_to_add], len(results) + 1):
                title = result.get('title', 'No title')
                link = result.get('link', 'No URL')
                snippet = result.get('snippet', 'No description')
                results.append(f"### 📰 News {i}: {title}\n**URL:** {link}\n**Summary:** {snippet}\n")
                st.write(f"✅ Found news: {title[:80]}...")

        if results:
            st.success(f"✅ Serper API found {len(results)} high-quality results")
            return "\n\n".join(results)
        else:
            st.warning("⚠️ No search results found via Serper API")
            return "❌ No search results found"

    except requests.exceptions.RequestException as e:
        error_msg = f"Serper API request failed: {str(e)}"
        st.error(f"❌ {error_msg}")
        return f"❌ {error_msg}"
    except Exception as e:
        error_msg = f"Unexpected error with Serper API: {str(e)}"
        st.error(f"❌ {error_msg}")
        return f"❌ {error_msg}"

@lru_cache(maxsize=100)
def cached_search(query: str, max_results: int, api_key: str) -> str:
    """Cache search results to avoid duplicate API calls"""
    return safe_serper_search(query, max_results, api_key)

# =====================================================================
# 🧠 ENHANCED GROQ LLM INTEGRATION WITH RATE LIMITING
# =====================================================================

class GroqLLM:
    """Custom Groq LLM wrapper"""

    def __init__(self, api_key, model, temperature=0.7):
        self.api_key = api_key
        self.model = model
        self.temperature = temperature
        self.base_url = "https://api.groq.com/openai/v1/chat/completions"

    def call(self, prompt, system_message=None):
        """Make API call to Groq"""
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }

        messages = []
        if system_message:
            messages.append({"role": "system", "content": system_message})
        messages.append({"role": "user", "content": prompt})

        payload = {
            "model": self.model,
            "messages": messages,
            "temperature": self.temperature,
            "max_tokens": 4000,
            "top_p": 1,
            "stream": False
        }

        try:
            response = requests.post(self.base_url, headers=headers, json=payload, timeout=60)
            response.raise_for_status()

            result = response.json()
            return result["choices"][0]["message"]["content"]

        except requests.exceptions.RequestException as e:
            st.error(f"❌ Groq API Request failed: {str(e)}")
            if hasattr(e, 'response') and e.response is not None:
                st.error(f"Response: {e.response.text}")
            return None
        except Exception as e:
            st.error(f"❌ Unexpected error: {str(e)}")
            return None

class RateLimitedGroqLLM:
    """Enhanced Groq LLM with rate limiting"""
    
    def __init__(self, api_key, model, temperature=0.7, requests_per_minute=30):
        self.llm = GroqLLM(api_key, model, temperature)
        self.requests_per_minute = requests_per_minute
        self.last_call_time = 0
        self.min_interval = 60.0 / requests_per_minute

    def call(self, prompt, system_message=None):
        # Rate limiting
        current_time = time.time()
        time_since_last_call = current_time - self.last_call_time
        
        if time_since_last_call < self.min_interval:
            sleep_time = self.min_interval - time_since_last_call
            with st.spinner(f"Rate limiting... waiting {sleep_time:.1f}s"):
                time.sleep(sleep_time)
        
        self.last_call_time = time.time()
        return self.llm.call(prompt, system_message)

# =====================================================================
# 📥 ENHANCED DOWNLOAD FUNCTIONS
# =====================================================================

def create_enhanced_downloads(linkedin_post, research_report, query):
    """Create multiple download formats with better formatting"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_filename = f"{query.replace(' ', '_')}_{timestamp}"
    
    # Enhanced TXT version
    txt_content = f"""
LINKEDIN POST
{'='*50}
{linkedin_post}

{'='*50}
RESEARCH REPORT
{'='*50}
{research_report}

Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    """
    txt_data = BytesIO(txt_content.encode('utf-8'))
    
    # Enhanced DOCX with better formatting
    doc = Document()
    
    # LinkedIn Post section
    doc.add_heading('LinkedIn Post', level=1)
    for paragraph in linkedin_post.split('\n'):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    
    doc.add_page_break()
    
    # Research Report section
    doc.add_heading('Research Report', level=1)
    for paragraph in research_report.split('\n'):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    
    # Add metadata
    doc.add_page_break()
    doc.add_heading('Document Information', level=2)
    doc.add_paragraph(f"Query: {query}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    docx_stream = BytesIO()
    doc.save(docx_stream)
    docx_stream.seek(0)
    
    return {
        'txt': (txt_data, f"{base_filename}.txt"),
        'docx': (docx_stream, f"{base_filename}.docx")
    }

# =====================================================================
# 🎯 ENHANCED RESEARCH WORKFLOW WITH PROGRESS TRACKING
# =====================================================================

def execute_research_workflow_with_progress(query, groq_llm, max_results, serper_key, post_tone, audience, include_stats, add_cta):
    """Enhanced workflow with progress tracking and customization"""
    
    progress_bar = st.progress(0)
    status_text = st.empty()
    
    # Step 1: Web search
    status_text.text("🔍 Searching the web...")
    search_results = cached_search(query, max_results, serper_key)
    progress_bar.progress(25)
    
    if search_results.startswith("❌"):
        st.error("Search failed. Please check your query and API key.")
        return None

    # Step 2: Research report
    status_text.text("📊 Analyzing results and generating research report...")
    
    research_prompt = f"""
    Analyze these search results and create a comprehensive research report about: "{query}"
    
    SEARCH RESULTS:
    {search_results}
    
    Current Date: {datetime.now().strftime('%Y-%m-%d')}
    
    Create a well-structured research report with:
    1. EXECUTIVE SUMMARY: Brief overview of key findings
    2. KEY FINDINGS: Main insights and discoveries
    3. DATA & STATISTICS: {"Include relevant data points and statistics" if include_stats else "Focus on qualitative insights"}
    4. IMPLICATIONS: What these findings mean in practice
    5. RECOMMENDATIONS: Actionable suggestions
    
    Target Audience: {', '.join(audience)}
    Tone: {post_tone.lower()}
    """
    
    research_report = groq_llm.call(research_prompt, "You are a professional research analyst.")
    progress_bar.progress(60)

    if not research_report:
        st.error("❌ Failed to generate research report")
        return None

    # Step 3: LinkedIn post
    status_text.text("💬 Creating LinkedIn post...")
    
    enhanced_linkedin_prompt = f"""
    {LINKEDIN_POST_PROMPT.format(research_content=research_report)}
    
    ADDITIONAL REQUIREMENTS:
    - Tone: {post_tone}
    - Target Audience: {', '.join(audience)}
    - {"Include relevant statistics and data" if include_stats else "Focus on conceptual insights"}
    - {"Add a clear call-to-action" if add_cta else "End with a thought-provoking question"}
    - Make it engaging for {post_tone.lower()} professional audience
    """
    
    linkedin_post = groq_llm.call(
        enhanced_linkedin_prompt,
        "You are a professional LinkedIn content creator with expertise in creating viral professional content. Your post must be of length 1,200 to 1,500 words"
    )
    progress_bar.progress(100)
    status_text.text("✅ Research complete!")
    
    return {
        "search_results": search_results,
        "research_report": research_report,
        "linkedin_post": linkedin_post
    }

# =====================================================================
# ⚙️ MAIN EXECUTION WITH SESSION STATE
# =====================================================================

# Initialize session state
if 'research_history' not in st.session_state:
    st.session_state.research_history = []

if 'last_results' not in st.session_state:
    st.session_state.last_results = None

def save_to_history(query, results):
    """Save research results to session history"""
    history_item = {
        'timestamp': datetime.now().isoformat(),
        'query': query,
        'linkedin_post': results['linkedin_post'],
        'research_report': results['research_report']
    }
    st.session_state.research_history.append(history_item)
    st.session_state.last_results = results

def main():
    """Main application logic."""
    
    query = st.text_area("🔎 Enter your research topic:", height=100, placeholder="e.g., The impact of AI on digital marketing in 2024...")

    final_groq_key = groq_api_key.strip() if groq_api_key else os.getenv("GROQ_API_KEY")
    final_serper_key = serper_api_key.strip() if serper_api_key else os.getenv("SERPER_API_KEY")

    # Validate API keys
    validation_errors = validate_api_keys(final_groq_key, final_serper_key)
    if validation_errors:
        for error in validation_errors:
            st.error(error)
        return

    # Initialize enhanced LLM with rate limiting
    groq_llm = RateLimitedGroqLLM(final_groq_key, model_options[selected_model], temperature)

    # Research execution
    col1, col2 = st.columns([1, 1])
    with col1:
        if st.button("🚀 Run Research", use_container_width=True, type="primary"):
            if not query.strip():
                st.warning("⚠️ Please enter a research topic.")
            elif len(query.strip()) < 5:
                st.warning("⚠️ Please enter a more detailed research topic (at least 5 characters).")
            else:
                result = execute_research_workflow_with_progress(
                    query, 
                    groq_llm, 
                    max_results, 
                    final_serper_key,
                    post_tone,
                    audience,
                    include_statistics,
                    add_call_to_action
                )
                if result:
                    save_to_history(query, result)
                    st.success("✅ Research complete!")
                    st.session_state.last_results = result

    with col2:
        if st.session_state.research_history:
            if st.button("📚 View Research History", use_container_width=True):
                st.session_state.show_history = True

    # Display results if available
    if st.session_state.last_results:
        result = st.session_state.last_results
        
        with st.expander("📋 LinkedIn Post", expanded=True):
            st.markdown(result["linkedin_post"])
            
            # Enhanced download buttons
            downloads = create_enhanced_downloads(
                result["linkedin_post"], 
                result["research_report"], 
                query
            )
            
            col1, col2 = st.columns(2)
            with col1:
                st.download_button(
                    "⬇️ Download as .TXT", 
                    downloads['txt'][0],
                    file_name=downloads['txt'][1],
                    mime="text/plain",
                    use_container_width=True
                )
            with col2:
                st.download_button(
                    "⬇️ Download as .DOCX", 
                    downloads['docx'][0],
                    file_name=downloads['docx'][1],
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        with st.expander("📊 Full Research Report", expanded=False):
            st.markdown(result["research_report"])

        with st.expander("🔍 Raw Search Results", expanded=False):
            st.markdown(result["search_results"])

    # Research history view
    if st.session_state.get('show_history', False) and st.session_state.research_history:
        st.markdown("---")
        st.subheader("📚 Research History")
        
        for i, item in enumerate(reversed(st.session_state.research_history[-5:]), 1):
            with st.expander(f"#{i}: {item['query']} - {datetime.fromisoformat(item['timestamp']).strftime('%Y-%m-%d %H:%M')}"):
                st.markdown("**LinkedIn Post:**")
                st.markdown(item['linkedin_post'][:500] + "..." if len(item['linkedin_post']) > 500 else item['linkedin_post'])
                
                if st.button(f"Load this result", key=f"load_{i}"):
                    st.session_state.last_results = item
                    st.rerun()

# =====================================================================
# 🧾 FOOTER
# =====================================================================

st.markdown("---")
st.caption("Built with ❤️ using Groq + Serper API | Reliable web search with AI-powered analysis")

if __name__ == "__main__":
    main()
